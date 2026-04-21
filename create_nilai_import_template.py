#!/usr/bin/env python3
"""
Script to create Excel template for importing student grades (Nilai)
Template format: Course info + student grades with multiple assessment components
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_nilai_import_template(output_path='Template_Import_Nilai.xlsx'):
    """Create an Excel template for importing student grades"""
    
    wb = Workbook()
    ws = wb.active
    ws.title = 'Data Nilai'
    
    # Set column widths
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 20
    ws.column_dimensions['H'].width = 20
    
    # Define styles
    title_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    title_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    info_font = Font(name='Calibri', size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Row 1: Kode Matakuliah
    ws['A1'] = 'Kode Matakuliah:'
    ws['A1'].font = info_font
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    ws['B1'] = ''  # User will fill this
    ws['B1'].font = info_font
    ws['B1'].border = Border(bottom=Side(style='thin'))
    
    # Row 2: Nama Matakuliah
    ws['A2'] = 'Nama Matakuliah:'
    ws['A2'].font = info_font
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center')
    ws['B2'] = ''  # User will fill this
    ws['B2'].font = info_font
    ws['B2'].border = Border(bottom=Side(style='thin'))
    
    # Row 3: Tahun Ajaran
    ws['A3'] = 'Tahun Ajaran:'
    ws['A3'].font = info_font
    ws['A3'].alignment = Alignment(horizontal='left', vertical='center')
    ws['B3'] = ''  # User will fill this (e.g., "2023/2024")
    ws['B3'].font = info_font
    ws['B3'].border = Border(bottom=Side(style='thin'))
    
    # Empty row
    # Row 5: Headers
    headers = ['NIM', 'Nama Mahasiswa', 'Aktivitas Partisipasi', 'Hasil Proyek', 'Kuis', 'Tugas', 'UTS', 'UAS']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    
    # Sample data rows (10 empty rows for user to fill)
    for row_num in range(6, 16):
        # NIM column
        cell_nim = ws.cell(row=row_num, column=1)
        cell_nim.alignment = Alignment(horizontal='center', vertical='center')
        cell_nim.border = border
        
        # Nama Mahasiswa column
        cell_nama = ws.cell(row=row_num, column=2)
        cell_nama.alignment = Alignment(horizontal='left', vertical='center')
        cell_nama.border = border
        
        # Score columns (C to H)
        for col_num in range(3, 9):
            cell_score = ws.cell(row=row_num, column=col_num)
            cell_score.alignment = Alignment(horizontal='center', vertical='center')
            cell_score.border = border
            # Format as number with 2 decimal places
            cell_score.number_format = '0.00'
    
    # Add instruction sheet
    instructions_ws = wb.create_sheet('Petunjuk')
    instructions_ws.column_dimensions['A'].width = 50
    instructions_ws.column_dimensions['B'].width = 50
    
    # Title
    title_cell = instructions_ws['A1']
    title_cell.value = 'PETUNJUK PENGGUNAAN TEMPLATE IMPORT NILAI'
    title_cell.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    # Instructions
    instructions = [
        ('', ''),
        ('1. Isi Informasi Matakuliah', 'Pada sheet "Data Nilai", isi kolom berikut:'),
        ('', '   - Kode Matakuliah: Contoh: PAFS6324'),
        ('', '   - Nama Matakuliah: Contoh: Termodinamika'),
        ('', '   - Tahun Ajaran: Contoh: 2020/2021'),
        ('', ''),
        ('2. Isi Data Mahasiswa', 'Mulai dari baris 6, isi:'),
        ('', '   - NIM: Nomor identitas mahasiswa (tidak boleh kosong)'),
        ('', '   - Nama Mahasiswa: Nama lengkap mahasiswa'),
        ('', ''),
        ('3. Isi Nilai', 'Masukkan nilai untuk setiap komponen penilaian:'),
        ('', '   - Aktivitas Partisipasi: Nilai 0-100'),
        ('', '   - Hasil Proyek: Nilai 0-100'),
        ('', '   - Kuis: Nilai 0-100'),
        ('', '   - Tugas: Nilai 0-100'),
        ('', '   - UTS: Nilai 0-100'),
        ('', '   - UAS: Nilai 0-100'),
        ('', ''),
        ('4. Format Data', 'Perhatikan hal-hal berikut:'),
        ('', '   - NIM harus unik (tidak ada duplikat)'),
        ('', '   - Nilai harus berupa angka (gunakan titik untuk desimal)'),
        ('', '   - Jika ada nilai yang kosong, gunakan 0 atau biarkan kosong'),
        ('', ''),
        ('5. Simpan File', 'Simpan file dengan format .xlsx sebelum di-upload'),
        ('', ''),
        ('Catatan:', 'Pastikan nama kolom tidak diubah agar proses import berjalan lancar'),
    ]
    
    for row_num, (col_a, col_b) in enumerate(instructions, 3):
        cell_a = instructions_ws.cell(row=row_num, column=1)
        cell_a.value = col_a
        cell_a.font = Font(name='Calibri', size=10, bold=(col_a and not col_a.startswith(' ')))
        cell_a.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        
        cell_b = instructions_ws.cell(row=row_num, column=2)
        cell_b.value = col_b
        cell_b.font = Font(name='Calibri', size=10)
        cell_b.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    
    # Save the workbook
    wb.save(output_path)
    print(f'✓ Template berhasil dibuat: {output_path}')
    return output_path

if __name__ == '__main__':
    # Create template in the Downloads folder
    import os
    downloads_path = os.path.expanduser('~/Downloads')
    template_path = os.path.join(downloads_path, 'Template_Import_Nilai.xlsx')
    create_nilai_import_template(template_path)
